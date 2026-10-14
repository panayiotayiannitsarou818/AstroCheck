"""v17.1: ποιότητα της αυτόματης διαδρομής — effort, γύρος ποιότητας, καλαίσθητο Word."""
from dataclasses import dataclass, field
from io import BytesIO
import types

import pytest
from docx import Document

import generator as G
from docx_builder import build_analysis_docx
from reference_loader import docx_text


# --- (α) επίπεδο προσπάθειας ------------------------------------------------
class _Stream:
    def __init__(self, calls, kwargs, fail):
        self.calls, self.kwargs, self.fail = calls, kwargs, fail

    def __enter__(self):
        self.calls.append(self.kwargs)
        if self.fail and "extra_body" in self.kwargs:
            import anthropic
            err = anthropic.BadRequestError.__new__(anthropic.BadRequestError)
            Exception.__init__(err, "output_config.effort: not supported for this model")
            raise err
        return self

    def __exit__(self, *a):
        return False

    def get_final_message(self):
        return types.SimpleNamespace(stop_reason="end_turn", content=[
            types.SimpleNamespace(type="thinking", text="(σκέψη)"),
            types.SimpleNamespace(type="text", text="ΚΕΙΜΕΝΟ"),
        ])


def _fake_anthropic(monkeypatch, fail=False):
    import anthropic
    calls = []

    class Client:
        def __init__(self, **kw):
            self.messages = types.SimpleNamespace(stream=lambda **kw: _Stream(calls, kw, fail))

    monkeypatch.setattr(anthropic, "Anthropic", Client)
    return calls


def test_claude_is_asked_for_high_effort_and_thinking_is_not_in_the_text(monkeypatch):
    calls = _fake_anthropic(monkeypatch)
    assert G.completer("Claude")("k", "εντολή", "claude-opus-5-5") == "ΚΕΙΜΕΝΟ"
    assert calls[0]["extra_body"] == {"output_config": {"effort": "high"}}


def test_model_without_effort_support_is_retried_without_it(monkeypatch):
    calls = _fake_anthropic(monkeypatch, fail=True)
    assert G.completer("Claude", "high")("k", "εντολή", "παλιό-μοντέλο") == "ΚΕΙΜΕΝΟ"
    assert "extra_body" in calls[0] and "extra_body" not in calls[1]


def test_effort_setting_is_parsed():
    assert G.effort_from("") == "high" and G.effort_from(None) == "high"
    assert G.effort_from("MAX") == "max" and G.effort_from("xhigh") == "xhigh"
    assert G.effort_from("default") == "" and G.effort_from("άκυρο") == "high"


# --- (β) γύρος ποιότητας μετά την έγκριση ---------------------------------
@dataclass
class R:
    ok: bool
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    def summary(self):
        return "ok" if self.ok else "σφάλματα"

    def details_lines(self):
        return list(self.errors)


def _validate(text):
    return R("ΛΑΘΟΣ" not in text, ["ΛΑΘΟΣ"] if "ΛΑΘΟΣ" in text else [], ["ΕΤΙΚΕΤΑ"] * text.count("ετικέτα"))


def _gen(answers):
    sent, it = [], iter(answers)
    out = G.generate_validated("k", "ΕΝΤΟΛΗ", _validate, complete=lambda k, t, m: sent.append(t) or next(it), polish=True)
    return out, sent


def test_quality_round_keeps_the_improved_text():
    out, sent = _gen(["κείμενο ετικέτα ετικέτα", "κείμενο καθαρό"])
    assert out.ok and out.text == "κείμενο καθαρό"
    assert out.rounds[-1].kind == "γύρος ποιότητας" and "ΕΤΙΚΕΤΑ" in sent[1]


def test_quality_round_never_loses_the_approval():
    out, _ = _gen(["κείμενο ετικέτα", "χαλασμένο ΛΑΘΟΣ"])
    assert out.ok and out.text == "κείμενο ετικέτα"


def test_no_quality_round_without_notes():
    out, sent = _gen(["καθαρό κείμενο"])
    assert out.ok and len(sent) == 1 and len(out.rounds) == 1


# --- (γ) καλαίσθητο Word από απλό κείμενο -----------------------------------
PLAIN = """Μεθοδολογία
Η ανάλυση βασίζεται σε ελεγμένα δεδομένα και ακολουθεί σταθερή σειρά σε κάθε Οίκο, με επεξήγηση των ορίων της.
1ος Οίκος — Ταυτότητα
Ήλιος–Σελήνη Εξάγωνο (orb 0°29′, Στενή/ισχυρή)
Πού βρίσκεται ο Οίκος
Η ακμή του Οίκου βρίσκεται στους Διδύμους και ο Οίκος εκτείνεται έως τον Καρκίνο, με έμφαση στην επικοινωνία.
Βασική δύναμη: α
Βασική πρόκληση: β
Κυβερνήτης: Ερμής
Τελικό συμπέρασμα: γ
Σύνθεση με τον υπόλοιπο χάρτη
Κείμενο σύνθεσης που συνεχίζει μετά το πλαίσιο και δεν πρέπει να μπει μέσα σε αυτό, γιατί είναι άλλη ενότητα.
2ος Οίκος — Αξίες
Κείμενο του δεύτερου Οίκου που είναι αρκετά μεγάλο ώστε να μοιάζει με κανονική παράγραφο κειμένου.
Κεντρικό θέμα: Η αξία μπορεί να στηρίζεται σε σαφείς συμφωνίες.
Τελική συνθετική εικόνα
Σύνθεση όλου του χάρτη σε μία παράγραφο κειμένου, με αρκετές λέξεις ώστε να είναι κανονική παράγραφος."""


def _doc(text):
    return Document(BytesIO(build_analysis_docx("A.M.", text)))


def test_plain_text_gets_real_headings_and_new_page_per_house():
    doc = _doc(PLAIN)
    h1 = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert [p.text for p in h1] == ["Μεθοδολογία", "1ος Οίκος — Ταυτότητα", "2ος Οίκος — Αξίες", "Τελική συνθετική εικόνα"]
    assert all(p.paragraph_format.page_break_before for p in h1[1:])
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert "Πού βρίσκεται ο Οίκος" in h2 and "Σύνθεση με τον υπόλοιπο χάρτη" in h2


def test_summary_box_holds_only_its_labels():
    doc = _doc(PLAIN)
    box = doc.tables[0].rows[0].cells[0].text
    assert "Κυβερνήτης: Ερμής" in box and "Κείμενο σύνθεσης" not in box


def test_client_document_title_and_central_theme_callout():
    doc = _doc(PLAIN)
    assert doc.paragraphs[0].text == "Ο γενέθλιος χάρτης σου"
    theme = next(p for p in doc.paragraphs if p.text.startswith("Κεντρικό θέμα:"))
    assert theme.runs[0].bold and theme._p.pPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr") is not None


def test_word_text_is_unchanged_by_the_formatting():
    text = docx_text(build_analysis_docx("A.M.", PLAIN))
    for line in PLAIN.splitlines():
        assert line in text, line
