"""Αγγλική ανάλυση: ο validator αναγνωρίζει αγγλικούς τίτλους, ονόματα,
ζώδια, όψεις και κατηγορίες βαρύτητας -- και πιάνει τα ίδια λάθη με την
ελληνική.

Οι πλήρεις αναλύσεις παράγονται συνθετικά από πραγματικά PDF-παραδείγματα.
Η ίδια γεννήτρια παράγει και ελληνική εκδοχή, ώστε να φαίνεται ότι οι δύο
γλώσσες ελέγχονται με τα ίδια κριτήρια.
"""

from pathlib import Path

import pytest

import lexicon_en as EN
import validator as V
from astrology import RULERS
from parser import parse_astrodienst_pdf
from prompts import build_master_prompt

FIXTURES = sorted((Path(__file__).parent / "fixtures").glob("astro_paradeigma*.pdf"))
FEMININE = {"Σελήνη", "Αφροδίτη"}


def _chart(path):
    return parse_astrodienst_pdf(path.read_bytes(), path.name)


@pytest.fixture(scope="module")
def chart2():
    return _chart(Path(__file__).parent / "fixtures" / "astro_paradeigma_2.pdf")


def build_analysis(chart, lang: str) -> str:
    """Πλήρης, σωστή ανάλυση: 12 Οίκοι, κυβερνήτες, πλαίσια σύνοψης, κάθε
    υποχρεωτική όψη στον Οίκο της, τελικές ενότητες και Παράρτημα."""
    english = lang == "en"
    name = EN.en if english else (lambda x: x)
    mandatory = V._mandatory_aspects(chart)
    out = []
    for n in range(1, 13):
        cusp = chart.cusps[n - 1]
        modern, traditional = RULERS[cusp.sign]
        rulers = [modern] + ([traditional] if traditional else [])
        if english:
            out.append(f"{EN.house_heading(n)} — Theme")
            out.append(f"The modern ruler is {name(modern)}.")
            if traditional:
                out.append(f"The traditional ruler is {name(traditional)}.")
        else:
            out.append(f"{n}ος Οίκος — Θέμα")
            art = "η" if modern in FEMININE else "ο"
            out.append(f"Κύριος κυβερνήτης είναι {art} {modern}.")
            if traditional:
                art = "η" if traditional in FEMININE else "ο"
                out.append(f"Παραδοσιακός κυβερνήτης είναι {art} {traditional}.")
        for p in chart.points:
            if p.house == n and p.kind in ("planet", "node"):
                if english:
                    out.append(f"{name(p.name)} is in the {EN.ordinal(n)} House.")
                    out.append(f"{name(p.name)} in {EN.en(p.sign)} gives colour here.")
                else:
                    out.append(f"Το σημείο {p.name} βρίσκεται στον {n}ο Οίκο.")
        involved = V._involved_points(chart, n)
        for a in mandatory:
            if a.first in involved or a.second in involved:
                aspect = name(a.aspect)
                out.append(
                    f"{name(a.first)}–{name(a.second)} {aspect} (orb {a.orb_text}, {name(a.weight)})."
                )
        out.append("")
        box = ", ".join(name(r) for r in rulers)
        if english:
            out += ["Key strength: text", "Key challenge: text", f"Ruler: {box}", "Final conclusion: text"]
        else:
            out += ["Βασική δύναμη: κείμενο", "Βασική πρόκληση: κείμενο", f"Κυβερνήτης: {box}", "Τελικό συμπέρασμα: κείμενο"]
        out.append("")
    titles = list(EN.SECTION_TITLES.values()) if english else list(EN.SECTION_TITLES)
    for title in titles[:3]:
        out += [title, "Text." if english else "Κείμενο.", ""]
    out.append(titles[3])
    for a in chart.aspects:
        out.append(f"{name(a.first)}–{name(a.second)} | {name(a.aspect)} | orb {a.orb_text} | {name(a.weight)}")
    return "\n".join(out)


def _house_line(text, needle):
    return next(line for line in text.splitlines() if needle in line)


# ---------------------------------------------------------------------------
# Σωστές αναλύσεις περνούν -- σε ΟΛΑ τα PDF-παραδείγματα, και στις δύο γλώσσες
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("path", FIXTURES, ids=lambda p: p.stem)
@pytest.mark.parametrize("lang", ["el", "en"])
def test_complete_analysis_passes(path, lang):
    chart = _chart(path)
    result = V.validate_analysis(chart, build_analysis(chart, lang))
    assert result.ok, result.details_lines()


def test_originally_reported_english_case_is_no_longer_rejected(chart2):
    """Το περιστατικό που εντοπίστηκε: «1st House», «Final synthesis»,
    «Appendix…» → «λείπουν οι 12 Οίκοι και όλες οι ενότητες»."""
    result = V.validate_analysis(chart2, build_analysis(chart2, "en"))
    assert not result.missing_houses and not result.missing_sections


# ---------------------------------------------------------------------------
# Τα ίδια λάθη εντοπίζονται και στα αγγλικά
# ---------------------------------------------------------------------------
def test_wrong_aspect_type_is_caught(chart2):
    text = build_analysis(chart2, "en")
    line = _house_line(text, "Venus–Saturn Square (orb 1°13′")
    bad = text.replace(line, line.replace("Square", "Trine"), 1)
    result = V.validate_analysis(chart2, bad)
    assert not result.ok
    assert any(a.first == "Αφροδίτη" and a.second == "Κρόνος" for a in result.wrong_aspect_type)


def test_wrong_weight_is_caught(chart2):
    text = build_analysis(chart2, "en")
    line = _house_line(text, "Venus–Saturn Square (orb 1°13′")
    bad = text.replace(line, line.replace("Tight/strong", "Standard"), 1)
    result = V.validate_analysis(chart2, bad)
    assert not result.ok
    assert any(a.first == "Αφροδίτη" and a.second == "Κρόνος" for a in result.wrong_weight)


def test_wrong_orb_is_caught(chart2):
    text = build_analysis(chart2, "en").replace("orb 1°13′", "orb 3°13′")
    result = V.validate_analysis(chart2, text)
    assert not result.ok


def test_missing_section_is_caught(chart2):
    text = build_analysis(chart2, "en").replace("Final synthesis\n", "Closing words\n")
    result = V.validate_analysis(chart2, text)
    assert result.missing_sections == ["Τελική συνθετική εικόνα"]


def test_missing_house_is_caught(chart2):
    text = build_analysis(chart2, "en")
    start = text.index("5th House")
    end = text.index("6th House")
    bad = text[:start] + text[end:]
    bad = bad.replace("in the 5th House", "")
    assert 5 in V.validate_analysis(chart2, bad).missing_houses


def test_inconsistent_ruler_box_is_caught(chart2):
    text = build_analysis(chart2, "en")
    seg = V._house_segments(text)[9]
    ruler_line = next(l for l in seg.splitlines() if l.startswith("Ruler:"))
    bad = text.replace(seg, seg.replace(ruler_line, "Ruler: Pluto"), 1)
    result = V.validate_analysis(chart2, bad)
    assert [h for h, _, _ in result.inconsistent_ruler_box] == [9]


@pytest.mark.parametrize(
    "text,expected",
    [
        ("The Moon in Scorpio can show intensity.", [("Σελήνη", "Σκορπιός", "Παρθένος")]),
        ("Mars, in Aries, shows drive.", [("Άρης", "Κριός", "Ζυγός")]),
        ("The Ascendant in Leo.", [("Ωροσκόπος", "Λέων", "Τοξότης")]),
    ],
)
def test_wrong_sign_is_caught(chart2, text, expected):
    assert [e[:3] for e in V._sign_claim_errors(chart2, text)] == expected


@pytest.mark.parametrize(
    "text",
    [
        "The Sun in Aquarius shows independence.",
        "Mars, in Libra, shows diplomacy.",
        "The North Node in Cancer is ruled by the Moon, which is in Virgo.",
        "Neptune, by degree, corresponds to Taurus.",
        "Saturn forms a square with Jupiter in Scorpio.",
        "Venus, ruler of the house whose cusp is in Taurus.",
        "Mercury encourages learning in Gemini-like curiosity.",
        "Venus, working in a Libra style, seeks balance.",
    ],
)
def test_non_placement_sign_phrases_do_not_trigger(chart2, text):
    assert V._sign_claim_errors(chart2, text) == []


def test_wrong_house_claim_is_caught(chart2):
    mars = next(p for p in chart2.points if p.name == "Άρης")
    wrong = mars.house % 12 + 1
    for text in (f"Mars is in the {EN.ordinal(wrong)} House.", f"Mars in the {EN.ordinal(wrong)} House brings drive."):
        errors = V._location_claim_errors(chart2, text)
        assert [(e[0], e[1], e[2]) for e in errors] == [("Άρης", wrong, mars.house)], text


@pytest.mark.parametrize(
    "text",
    [
        "Mars is the ruler of the 7th House.",
        "Mars rules the 7th House and the 12th House.",
        "Mars, in Libra, connects with the themes of the 7th House.",
    ],
)
def test_ruler_or_thematic_link_is_not_a_placement(chart2, text):
    assert V._location_claim_errors(chart2, text) == []


def test_thematic_placement_is_caught(chart2):
    mars = next(p for p in chart2.points if p.name == "Άρης")
    if mars.house == 7:
        pytest.skip("Ο Άρης βρίσκεται ήδη στον 7ο Οίκο σε αυτό το παράδειγμα.")
    errors = V._location_claim_errors(chart2, "Mars is in the field of relationships and partnership.")
    assert [(e[0], e[1]) for e in errors] == [("Άρης", 7)]


def test_mirrored_axis_type_is_ignored_in_english():
    fragment = "Moon–Midheaven trine, which activates, as a sextile, the Imum Coeli."
    start = fragment.index("sextile")
    assert V._is_mirrored_axis_type(fragment, start)


@pytest.mark.parametrize(
    "section,heading",
    [
        ("Τελική συνθετική εικόνα", "Final synthesis"),
        ("Τελική συνθετική εικόνα", "## Overall picture"),
        ("Συμβολική κατεύθυνση εξέλιξης", "The Lunar Nodes"),
        ("Συμβολική κατεύθυνση εξέλιξης", "North and South Node"),
        ("Προτάσεις προσωπικής ανάπτυξης", "3. Suggestions for personal growth and support"),
        ("Προτάσεις προσωπικής ανάπτυξης", "Practical suggestions for growth"),
        ("Παράρτημα επιβεβαιωμένων όψεων", "Technical appendix"),
        ("Παράρτημα επιβεβαιωμένων όψεων", "Table of confirmed aspects"),
    ],
)
def test_english_equivalent_headings(section, heading):
    assert V._has_section(f"Intro.\n{heading}\nSection text.", section)


def test_english_mention_in_prose_does_not_count_as_section():
    text = "Text.\nThe technical appendix of your report is available on request.\nEnd."
    assert not V._has_section(text, "Παράρτημα επιβεβαιωμένων όψεων")


# ---------------------------------------------------------------------------
# Έντυπο πελάτη στα αγγλικά
# ---------------------------------------------------------------------------
def _en_source():
    houses = "\n".join(f"{EN.house_heading(n)}\nText." for n in range(1, 13))
    titles = list(EN.SECTION_TITLES.values())
    return houses + "\n" + "\n".join(f"{t}\nText." for t in titles)


def _en_rewrite(extra="", closing="", chart=None):
    from sample_texts import houses_en, section_body

    if chart is None:
        chart = _chart(Path(__file__).parent / "fixtures" / "astro_paradeigma_2.pdf")
    houses = houses_en(chart, {2: extra})
    titles = list(EN.SECTION_TITLES.values())[:3]
    return houses + "\n" + "\n".join(f"{t}\n{section_body('en')}" for t in titles) + "\n" + closing


def test_english_client_document_passes(chart2):
    r = V.validate_rewrite(chart2, _en_source(), _en_rewrite())
    assert r.ok, r.details_lines()


def test_english_client_document_rejects_technical_data(chart2):
    r = V.validate_rewrite(
        chart2, _en_source(), _en_rewrite(" Venus at 23°21′ has an orb of 1°13′, Tight/strong.")
    )
    cats = {c for c, _ in r.technical_data}
    assert not r.ok and {"μοίρες/λεπτά", "orb", "κατηγορία βαρύτητας"} <= cats


def test_english_client_document_rejects_appendix(chart2):
    text = _en_rewrite() + "\nAppendix of confirmed aspects"
    r = V.validate_rewrite(chart2, _en_source(), text)
    assert any(c == "Παράρτημα/πίνακας όψεων" for c, _ in r.technical_data)


def test_english_mathematical_check_sentence_is_flagged(chart2):
    old = (
        "This analysis is based on a full mathematical check of all the astrological "
        "aspects in your chart; the technical details are available if you would like to see them."
    )
    r = V.validate_rewrite(chart2, _en_source(), _en_rewrite(closing=old))
    assert not r.ok
    assert any(c.startswith("πρόταση μαθηματικού ελέγχου") for c, _ in r.technical_data)


def test_english_client_document_catches_missing_section(chart2):
    text = _en_rewrite().replace("Symbolic direction of growth\n", "")
    r = V.validate_rewrite(chart2, _en_source(), text)
    assert r.missing_source_sections == ["Συμβολική κατεύθυνση εξέλιξης"]


def test_english_personal_claims_are_caught():
    found = V._unauthorized_personal_claims({}, "As a teacher, you have two children to care for.")
    assert {c for c, _ in found} == {"επάγγελμα/σπουδές", "οικογενειακή κατάσταση"}


# ---------------------------------------------------------------------------
# Η εντολή δίνει στο μοντέλο ΤΟΥΣ ΙΔΙΟΥΣ όρους που ελέγχει ο validator
# ---------------------------------------------------------------------------
def test_english_prompt_contains_terminology_and_translated_appendix(chart2):
    prompt = build_master_prompt(chart2, {}, "Αγγλικά", "ΟΔΗΓΙΕΣ", "ΥΦΟΣ")
    assert "ENGLISH OUTPUT — MANDATORY TERMINOLOGY" in prompt
    for a in chart2.aspects:
        line = f"{EN.en(a.first)}–{EN.en(a.second)} | {EN.en(a.aspect)} | orb {a.orb_text} | {EN.en(a.weight)}"
        assert line in prompt
    for title in EN.SECTION_TITLES.values():
        assert title in prompt


def test_greek_prompt_is_unchanged(chart2):
    prompt = build_master_prompt(chart2, {}, "Ελληνικά", "ΟΔΗΓΙΕΣ", "ΥΦΟΣ")
    assert "ENGLISH OUTPUT" not in prompt


def test_english_word_gets_english_title_and_summary_box(chart2):
    from io import BytesIO

    from docx import Document

    from docx_builder import build_analysis_docx

    doc = Document(BytesIO(build_analysis_docx("Client", build_analysis(chart2, "en"))))
    assert doc.paragraphs[0].text == "Complete Astrological Analysis"
    assert doc.tables, "τα πλαίσια σύνοψης πρέπει να γίνονται πίνακες"
    first = doc.tables[0].rows[0].cells[0].paragraphs[0]
    assert first.text.startswith("Key strength: ")
    assert any(r.bold and r.text == "Key strength: " for r in first.runs)
